"""Regenerate tests/fixtures/program_mini.xlsx (dev only, needs openpyxl).

A small synthetic program with the same sheet layout as the real
Program_Marathon_<NAME>.xlsx files. Contains no personal data.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

GREY = PatternFill("solid", fgColor="FFF2F2F2")
ORANGE = PatternFill("solid", fgColor="FFFCE4D6")

wb = openpyxl.Workbook()
wb.active.title = "Panduan"
wb["Panduan"].append(["PROGRAM MARATHON — UJI", None, None])
wb["Panduan"].append([None, "Race", "Minggu 13 Juni 2027"])

ws = wb.create_sheet("Kalender")
ws.append(["KALENDER HARIAN UJI — MS7 s/d MS10"])
ws.append(["Baris abu-abu = minggu ringan."])
ws.append(["MS", "TANGGAL MULAI", "FASE", "SENIN\n(ISTIRAHAT)", "SELASA", "RABU", "KAMIS", "JUMAT",
           "SABTU\n(LONG RUN / LOMBA)", "MINGGU", "KM TOTAL", "LONG RUN km", "CATATAN COACHING"])
rows = [(7, "03 Agu 2026", "1 Base", 28, 14, "Mulai", None),
        (8, "10 Agu 2026", "1 Base", 24, 12, "Minggu ringan", GREY),
        (9, "17 Agu 2026", "2 Support", 21, "-", "LOMBA 10K", ORANGE),
        (10, "24 Agu 2026", "2 Support", 35, 20, "DIUBAH. Bukan minggu ringan lagi", GREY)]
for ms, date, phase, km, lr, note, fill in rows:
    ws.append([ms, date, phase, "", "", "", "", "", "", "", km, lr, note])
    if fill:
        ws.cell(ws.max_row, 1).fill = fill

cp = wb.create_sheet("Checkpoint")
cp.append(["JADWAL LOMBA"])
cp.append(["MS", "Tanggal", "Acara", "Fungsi", "Sesudahnya"])
cp.append([9, "Sab 22 Agu 2026", "LOMBA 10K", "", ""])
cp.append([51, "MINGGU 13 JUN 2027", "MARATON", "Lomba utama", "-"])
cp.append(["TARGET WAKTU MARATON"])
cp.append(["Tingkat", "Waktu", "Pace", "Peluang"])
cp.append(["A - Ambisius", "di bawah 4:45:00", "", "30%"])
cp.append(["B - TARGET UTAMA", "4:45:00 - 5:15:00", "", "50%"])
cp.append(["C - Aman", "5:30 - 5:45", "", "-"])
cp.append(["SYARATNYA", "..."])

sl = wb.create_sheet("Tidur dan Pemulihan")
sl.append(["Target minimum", "6 jam 30 menit"])
sl.append(["Target malam sebelum long run", "7 jam"])

lb = wb.create_sheet("Log Berat")
lb.append(["LOG BERAT"])
lb.append(["MS", "Tanggal Minggu", "Target kg", "Aktual kg"])
lb.append([7, "09 Agu 2026", 81.3, None])
lb.append([8, "16 Agu 2026", 81.0, None])
lb.append([7, "09 Agu 2026", 99.0, 80.0])  # example row at the end: must be ignored

wb.save(Path(__file__).parent / "fixtures" / "program_mini.xlsx")
